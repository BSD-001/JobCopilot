"""读取文本、PDF 和 Word 文档。"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx"}


def read_document(file_path: str | Path) -> str:
    """根据文件后缀读取文档并返回纯文本。"""

    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在：{path}")
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError(
            f"暂不支持 {path.suffix} 文件。"
            f"支持的文件类型：{', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )

    return read_document_bytes(path.read_bytes(), path.name)


def read_document_bytes(content: bytes, file_name: str) -> str:
    """读取浏览器上传文件的字节内容并返回纯文本。"""

    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(
            f"暂不支持 {suffix} 文件。"
            f"支持的文件类型：{', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )
    if suffix in {".txt", ".md"}:
        return content.decode("utf-8-sig")
    if suffix == ".pdf":
        return _read_pdf(BytesIO(content))
    return _read_docx(BytesIO(content))


def _read_pdf(source) -> str:
    reader = PdfReader(source)
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(page for page in pages if page)


def _read_docx(source) -> str:
    document = Document(source)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cell for cell in cells if cell))

    clean_parts = [part for part in parts if part]
    known_parts = set(clean_parts)
    for textbox_text in _iter_textbox_text(document):
        if textbox_text not in known_parts:
            clean_parts.append(textbox_text)
            known_parts.add(textbox_text)

    return "\n".join(clean_parts)


def _iter_textbox_text(document) -> list[str]:
    """提取 Word 文本框中的文字。"""

    results: list[str] = []
    for container in document.element.body.iter(qn("w:txbxContent")):
        for paragraph in container.iter(qn("w:p")):
            text = "".join(
                node.text or "" for node in paragraph.iter(qn("w:t"))
            ).strip()
            if text:
                results.append(text)
    return results

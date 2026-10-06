"""文档读取功能测试。"""

from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from core.document_reader import read_document, read_document_bytes


class DocumentReaderTests(unittest.TestCase):
    def test_reads_plain_text(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "resume.txt"
            path.write_text("教育背景：人工智能专业", encoding="utf-8")
            self.assertIn("人工智能", read_document(path))

    def test_reads_word_paragraph_and_table(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "resume.docx"
            document = Document()
            document.add_paragraph("项目经历：人工智能宣传项目")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "技能"
            table.cell(0, 1).text = "Python"
            document.save(path)

            text = read_document(path)
            self.assertIn("人工智能宣传项目", text)
            self.assertIn("Python", text)

    def test_reads_uploaded_text_bytes(self) -> None:
        text = read_document_bytes(
            "教育背景：人工智能专业".encode("utf-8"),
            "resume.txt",
        )
        self.assertIn("人工智能", text)

    def test_rejects_unsupported_file_type(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "resume.csv"
            path.write_text("a,b", encoding="utf-8")
            with self.assertRaises(ValueError):
                read_document(path)


if __name__ == "__main__":
    unittest.main()

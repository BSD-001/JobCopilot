"""接口闭环与文件隔离测试，不请求真实模型、不删除用户记录。"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
from time import monotonic, sleep
from unittest import TestCase
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from openai import APITimeoutError
from openai import OpenAIError
from pypdf import PdfWriter

from core.history_store import save_history_record
from tests.web_fixtures import SAMPLE_JD, SAMPLE_RESUME, WebFakeClient
from webapp import create_app


class WebAppTests(TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.root = Path(self.temporary.name) / "history"
        self.app = create_app(self.root, WebFakeClient)
        self.client = TestClient(self.app)

    def tearDown(self):
        self.app.state.executor.shutdown(wait=True)
        self.client.close()
        self.temporary.cleanup()

    def submit(self, **overrides):
        data = {"api_key": "test-key-not-a-real-secret", "resume_mode": "text", "jd_mode": "text", "resume_text": SAMPLE_RESUME, "jd_text": SAMPLE_JD}
        data.update(overrides)
        return self.client.post("/api/analyses", data=data)

    def wait_job(self, job_id):
        deadline = monotonic() + 10
        while monotonic() < deadline:
            response = self.client.get(f"/api/analyses/{job_id}")
            self.assertEqual(response.status_code, 200)
            if response.json()["status"] in {"completed", "failed"}:
                return response.json()
            sleep(0.01)
        self.fail("离线任务未结束")

    def test_index_and_static_files(self):
        self.assertEqual(self.client.get("/").status_code, 200)
        self.assertIn("让每一次求职", self.client.get("/").text)
        self.assertEqual(self.client.get("/static/app.js").status_code, 200)

    def test_missing_inputs_are_rejected_individually(self):
        for field in ("api_key", "resume_text", "jd_text"):
            with self.subTest(field=field):
                self.assertEqual(self.submit(**{field: ""}).status_code, 400)

    def test_text_analysis_history_download_and_restart(self):
        response = self.submit()
        self.assertEqual(response.status_code, 202)
        job_id = response.json()["job_id"]
        result = self.wait_job(job_id)
        self.assertEqual(result["status"], "completed", result)
        self.assertFalse(result["has_tailored_resume"])
        self.assertEqual(len(result["analysis_data"]["interview"]["questions"]), 10)
        self.assertEqual(len(result["analysis_data"]["learning_plan"]["weeks"]), 4)
        self.assertNotIn("test-key-not-a-real-secret", str(result))
        self.assertIn("【文档交付】", result["report_text"])
        self.assertNotIn("模型扩写内容", result["report_text"])
        record_id = result["record_id"]
        self.assertTrue((self.root / record_id / "analysis.json").exists())
        self.assertEqual(len(self.client.get("/api/history").json()), 1)
        download = self.client.get(f"/api/analyses/{job_id}/downloads/report")
        self.assertEqual(download.status_code, 200)
        self.assertIn("filename*=UTF-8", download.headers["content-disposition"])
        self.assertEqual(self.client.get(f"/api/analyses/{job_id}/downloads/resume").status_code, 404)
        fresh_app = create_app(self.root, WebFakeClient)
        with TestClient(fresh_app) as fresh:
            recovered = fresh.get(f"/api/history/{record_id}")
            self.assertEqual(recovered.status_code, 200)
            self.assertEqual(recovered.json()["analysis_data"], result["analysis_data"])
        fresh_app.state.executor.shutdown(wait=True)

    def test_docx_download_preserves_template(self):
        template = Document()
        template.add_paragraph("示例求职者").runs[0].bold = True
        template.add_paragraph("制作并交付3份项目说明。")
        output = BytesIO()
        template.save(output)
        response = self.client.post("/api/analyses", data={"api_key": "mock", "resume_mode": "file", "jd_mode": "text", "jd_text": SAMPLE_JD}, files={"resume_file": ("resume.docx", output.getvalue())})
        result = self.wait_job(response.json()["job_id"])
        self.assertTrue(result["has_tailored_resume"], result)
        download = self.client.get(f"/api/history/{result['record_id']}/downloads/resume")
        document = Document(BytesIO(download.content))
        self.assertTrue(document.paragraphs[0].runs[0].bold)
        self.assertIn("【文档交付】", document.paragraphs[1].text)
        self.assertIn("3份", document.paragraphs[1].text)
        self.assertNotIn("待补充", document.paragraphs[1].text)

    def test_legacy_record_is_readable_without_rewrite(self):
        record = save_history_record("旧岗位", "# 旧报告", history_root=self.root)
        path = self.root / record["record_id"] / "report.md"
        before = path.read_bytes()
        result = self.client.get(f"/api/history/{record['record_id']}")
        self.assertEqual(result.status_code, 200)
        self.assertIsNone(result.json()["analysis_data"])
        self.assertIn("<h1>", result.json()["report_html"])
        self.assertEqual(path.read_bytes(), before)
        self.assertFalse((path.parent / "analysis.json").exists())

    def test_existing_placeholder_blocks_word_download(self):
        template = Document()
        template.add_paragraph("示例求职者 [待补充：联系方式]")
        template.add_paragraph("制作并交付3份项目说明。")
        output = BytesIO()
        template.save(output)
        response = self.client.post("/api/analyses", data={"api_key": "mock", "resume_mode": "file", "jd_text": SAMPLE_JD}, files={"resume_file": ("resume.docx", output.getvalue())})
        result = self.wait_job(response.json()["job_id"])
        self.assertEqual(result["status"], "completed")
        self.assertFalse(result["has_tailored_resume"])
        self.assertTrue(any("待补充标记" in item for item in result["warnings"]))

    def test_delete_only_one_record_and_blocks_traversal(self):
        first = save_history_record("第一条", "报告", history_root=self.root)
        second = save_history_record("第二条", "报告", history_root=self.root)
        outside = self.root.parent / "keep.txt"
        outside.write_text("保留", encoding="utf-8")
        response = self.client.delete(f"/api/history/{first['record_id']}")
        self.assertEqual(response.status_code, 200)
        self.assertFalse((self.root / first["record_id"]).exists())
        self.assertTrue((self.root / second["record_id"]).exists())
        self.assertTrue(outside.exists())
        self.assertEqual(self.client.delete("/api/history/invalid").status_code, 400)

    def test_upload_validation_and_modes(self):
        for name, content in (("resume.csv", b"data"), ("resume.txt", b"")):
            response = self.client.post("/api/analyses", data={"api_key": "mock", "jd_text": SAMPLE_JD}, files={"resume_file": (name, content)})
            self.assertEqual(response.status_code, 400)
        with patch("webapp.MAX_UPLOAD_SIZE", 4):
            response = self.client.post("/api/analyses", data={"api_key": "mock", "jd_text": SAMPLE_JD}, files={"resume_file": ("resume.txt", b"12345")})
            self.assertEqual(response.status_code, 413)
        self.assertEqual(self.submit(resume_mode="unknown").status_code, 400)

    def test_upload_jd_and_paste_resume(self):
        response = self.client.post("/api/analyses", data={"api_key": "mock", "resume_mode": "text", "resume_text": SAMPLE_RESUME, "jd_mode": "file"}, files={"jd_file": ("job.md", SAMPLE_JD.encode("utf-8"))})
        self.assertEqual(self.wait_job(response.json()["job_id"])["status"], "completed")

    def test_unreadable_or_empty_document_allows_retry(self):
        pdf = PdfWriter()
        pdf.add_blank_page(width=100, height=100)
        content = BytesIO()
        pdf.write(content)
        for filename, data in (("blank.pdf", content.getvalue()), ("blank.txt", b"   \n"), ("broken.docx", b"not-a-document")):
            response = self.client.post("/api/analyses", data={"api_key": "mock", "jd_text": SAMPLE_JD}, files={"resume_file": (filename, data)})
            result = self.wait_job(response.json()["job_id"])
            self.assertEqual(result["status"], "failed")
            self.assertTrue(result["error"])
            self.assertIsNone(self.app.state.active_job)
            if filename.endswith(".pdf"):
                self.assertIn("扫描型PDF", result["error"])
        self.assertEqual(self.wait_job(self.submit().json()["job_id"])["status"], "completed")

    def test_model_failure_hides_key_and_releases_task(self):
        class FailedClient(WebFakeClient):
            def generate_json(self, **kwargs):
                raise OpenAIError("test-key-not-a-real-secret")
        failed_app = create_app(self.root, FailedClient)
        try:
            with TestClient(failed_app) as client:
                data = {"api_key": "test-key-not-a-real-secret", "resume_mode": "text", "resume_text": SAMPLE_RESUME, "jd_text": SAMPLE_JD}
                job_id = client.post("/api/analyses", data=data).json()["job_id"]
                deadline = monotonic() + 5
                while monotonic() < deadline:
                    result = client.get(f"/api/analyses/{job_id}").json()
                    if result["status"] == "failed":
                        break
                    sleep(0.01)
                self.assertEqual(result["status"], "failed")
                self.assertNotIn(data["api_key"], str(result))
                self.assertIsNone(failed_app.state.active_job)
        finally:
            failed_app.state.executor.shutdown(wait=True)

    def test_unexpected_error_has_safe_stage_and_location(self):
        secret = "test-key-not-a-real-secret"
        private_text = "不应记录的简历正文"

        def fail_pipeline(**kwargs):
            kwargs["on_progress"]("整理分析结果")
            raise TypeError(f"{secret} {private_text}")

        with self.assertLogs("webapp", level="ERROR") as logs:
            with patch("webapp.run_pipeline_result", side_effect=fail_pipeline):
                result = self.wait_job(self.submit().json()["job_id"])
        self.assertEqual(result["status"], "failed")
        self.assertIn("整理分析结果", result["error"])
        self.assertIn("TypeError", result["error"])
        self.assertEqual(result["error_details"]["stage"], "整理分析结果")
        self.assertEqual(result["error_details"]["exception_type"], "TypeError")
        self.assertTrue(result["error_details"]["frames"])
        self.assertNotIn(secret, str(result) + str(logs.output))
        self.assertNotIn(private_text, str(result) + str(logs.output))
        self.assertIsNone(self.app.state.active_job)

    def test_history_failure_keeps_download(self):
        with patch("webapp.save_history_record", side_effect=OSError("disk")):
            result = self.wait_job(self.submit().json()["job_id"])
        self.assertEqual(result["status"], "completed")
        self.assertTrue(result["warnings"])
        job_id = next(iter(self.app.state.jobs))
        self.assertEqual(self.client.get(f"/api/analyses/{job_id}/downloads/report").status_code, 200)

    def test_busy_and_timeout_are_recoverable(self):
        entered = Event()
        release = Event()

        class BlockingClient(WebFakeClient):
            def generate_json(self, *args, **kwargs):
                entered.set()
                release.wait(5)
                raise APITimeoutError(request=None)

        blocking_app = create_app(self.root, BlockingClient)
        try:
            with TestClient(blocking_app) as client:
                data = {"api_key": "mock", "resume_mode": "text", "resume_text": SAMPLE_RESUME, "jd_text": SAMPLE_JD}
                response = client.post("/api/analyses", data=data)
                self.assertTrue(entered.wait(2))
                job_id = response.json()["job_id"]
                self.assertEqual(client.get(f"/api/analyses/{job_id}").json()["stage"], "解析简历")
                self.assertEqual(client.get(f"/api/analyses/{job_id}/downloads/report").status_code, 409)
                self.assertEqual(client.post("/api/analyses", data=data).status_code, 409)
                release.set()
                deadline = monotonic() + 5
                while monotonic() < deadline:
                    result = client.get(f"/api/analyses/{response.json()['job_id']}").json()
                    if result["status"] == "failed":
                        break
                    sleep(0.01)
                self.assertEqual(result["status"], "failed")
                self.assertIn("超时", result["error"])
        finally:
            release.set()
            blocking_app.state.executor.shutdown(wait=True)

    def test_external_origin_and_html_are_blocked(self):
        self.assertEqual(self.client.get("/api/history", headers={"Origin": "https://example.org"}).status_code, 403)
        record = save_history_record("示例", '<script>alert(1)</script><a href="javascript:alert(1)">危险</a>', history_root=self.root)
        html = self.client.get(f"/api/history/{record['record_id']}").json()["report_html"]
        self.assertNotIn("<script", html)
        self.assertNotIn("javascript:", html)

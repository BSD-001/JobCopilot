"""完整流程的离线测试。"""

from __future__ import annotations

import unittest
import re
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import main as main_module
from docx import Document


class FakeDeepSeekClient:
    def generate_json(self, system_prompt: str, user_prompt: str, **kwargs):
        if "请解析简历文本" in user_prompt:
            return {
                "education": [{"school": "示例大学"}],
                "projects": [{"name": "示例项目"}],
                "skills": ["Python"],
                "experiences": [],
            }
        if "请解析岗位描述" in user_prompt:
            return {
                "job_title": "内容运营",
                "responsibilities": ["内容策划"],
                "requirements": ["沟通能力"],
                "keywords": ["内容运营"],
            }
        if "匹配度评分" in user_prompt:
            return {
                "score": 80,
                "matched_points": ["具备内容策划经历"],
                "missing_points": ["数据分析"],
                "reason": "基础匹配。",
            }
        if "重写简历" in user_prompt:
            return {
                "optimized_projects": [
                    {"title": "项目一", "content": "优化内容一"},
                    {"title": "项目二", "content": "优化内容二"},
                    {"title": "项目三", "content": "优化内容三"},
                ]
            }
        if "面试问题" in user_prompt:
            return {
                "questions": [
                    {"question": "问题", "answer": "回答", "focus": "重点"}
                ]
            }
        if "resume_blocks" in user_prompt:
            block_ids = [
                int(value)
                for value in re.findall(r'"block_id":\s*(\d+)', user_prompt)
            ][:2]
            return {
                "target_keywords": ["内容运营"],
                "summary": "针对岗位调整项目表达。",
                "replacements": [
                    {
                        "block_id": block_ids[1],
                        "original_text": "",
                        "new_text": "针对岗位强化后的项目经历",
                        "reason": "突出内容运营关键词。",
                    }
                ],
                "removals": [
                    {
                        "block_id": block_ids[0],
                        "original_text": "无关校园经历",
                        "reason": "与目标岗位关联较弱。",
                    }
                ],
                "skill_boosts": [
                    {
                        "skill": "数据分析",
                        "suggested_wording": "正在学习数据分析基础",
                        "reason": "岗位高频要求。",
                    }
                ],
            }
        if "生成一段可直接发送给招聘平台 HR" in user_prompt:
            return {
                "greeting": "您好，我应聘内容运营。我有内容策划经验，方便看下简历吗？"
            }
        return {
            "weeks": [
                {
                    "week": 1,
                    "goal": "目标",
                    "tasks": ["任务"],
                    "deliverable": "成果",
                }
            ],
            "summary": "总结",
        }

    def generate_text(self, system_prompt: str, user_prompt: str, **kwargs):
        return (
            "# 求职分析报告\n\n"
            "## 匹配度分析\n\n"
            "## 简历优化建议\n\n"
            "## 面试预测题\n\n"
            "## 三十天学习计划\n\n"
            "## 项目成果摘要\n\n"
            "流水线测试成功。\n\n"
            "## 线上投递打招呼语\n\n"
            "这段内容会被短招呼语替换。"
        )


class PipelineTests(unittest.TestCase):
    def test_pipeline_result_returns_job_title(self) -> None:
        result = main_module.run_pipeline_result(
            resume_text="教育背景：人工智能专业",
            job_description_text="岗位职责：负责内容策划",
            llm_client=FakeDeepSeekClient(),
        )

        self.assertEqual(result.job_title, "内容运营")
        self.assertIn("求职分析报告", result.report_text)

    def test_pipeline_returns_markdown_report(self) -> None:
        report = main_module.run_pipeline(
            resume_text="教育背景：人工智能专业",
            job_description_text="岗位职责：负责内容策划",
            llm_client=FakeDeepSeekClient(),
        )
        self.assertIn("求职分析报告", report)
        self.assertIn("流水线测试成功", report)
        for heading in (
            "匹配度分析",
            "简历优化建议",
            "面试预测题",
            "三十天学习计划",
            "项目成果摘要",
            "线上投递打招呼语",
        ):
            self.assertIn(heading, report)
        greeting = report.split("## 线上投递打招呼语", 1)[1].strip()
        self.assertEqual(
            greeting,
            "您好，我应聘内容运营。我有内容策划经验，方便看下简历吗？",
        )
        self.assertLessEqual(len(greeting.replace(" ", "")), 65)

    def test_command_line_writes_report_file(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            resume_path = directory / "resume.txt"
            job_description_path = directory / "job_description.txt"
            output_path = directory / "report.md"
            resume_path.write_text("教育背景：人工智能专业", encoding="utf-8")
            job_description_path.write_text(
                "岗位职责：负责内容策划",
                encoding="utf-8",
            )

            with (
                patch.object(main_module, "DeepSeekClient", FakeDeepSeekClient),
                patch.object(main_module, "OUTPUT_PATH", output_path),
            ):
                exit_code = main_module.main(
                    [
                        "--resume",
                        str(resume_path),
                        "--jd",
                        str(job_description_path),
                    ]
                )

            self.assertEqual(exit_code, 0)
            self.assertTrue(output_path.exists())
            self.assertIn("匹配度分析", output_path.read_text(encoding="utf-8"))

    def test_pipeline_writes_tailored_word_resume(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            resume_path = directory / "resume.docx"
            tailored_path = directory / "tailored_resume.docx"
            job_description_path = directory / "job_description.txt"
            document = Document()
            document.add_paragraph("无关校园经历")
            document.add_paragraph("原项目经历")
            document.save(resume_path)
            job_description_path.write_text(
                "岗位职责：负责内容策划",
                encoding="utf-8",
            )

            main_module.run_pipeline(
                resume_text="原项目经历",
                job_description_text="岗位职责：负责内容策划",
                llm_client=FakeDeepSeekClient(),
                resume_file_path=resume_path,
                tailored_resume_output=tailored_path,
            )

            self.assertTrue(tailored_path.exists())
            rewritten = Document(tailored_path)
            texts = [paragraph.text for paragraph in rewritten.paragraphs]
            self.assertIn("针对岗位强化后的项目经历", texts)
            self.assertNotIn("无关校园经历", texts)


if __name__ == "__main__":
    unittest.main()

"""岗位定制简历的模板提取和回填测试。"""

from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from agents.tailored_resume_agent import TailoredResumeAgent
from core.resume_template import apply_paragraph_replacements, extract_paragraph_blocks


class FakeDeepSeekClient:
    def generate_json(self, system_prompt: str, user_prompt: str, **kwargs):
        return {
            "target_keywords": ["内容运营", "项目管理"],
            "summary": "针对岗位强化内容运营和项目管理表达。",
            "replacements": [
                {
                    "block_id": 1,
                    "original_text": "原项目经历",
                    "new_text": "围绕内容运营目标完成项目方案设计、协作推进和结果复盘。",
                    "reason": "优先呈现岗位关注的运营与项目推进能力。",
                }
            ],
            "removals": [
                {
                    "block_id": 2,
                    "original_text": "无关校园经历",
                    "reason": "与目标岗位关联较弱。",
                }
            ],
            "skill_boosts": [
                {
                    "skill": "数据分析",
                    "suggested_wording": "正在学习数据分析基础",
                    "reason": "岗位高频要求。",
                }
            ],
        }


class TailoredResumeTests(unittest.TestCase):
    def test_extract_and_apply_paragraph_replacement(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            source = directory / "resume.docx"
            output = directory / "tailored.docx"

            document = Document()
            paragraph = document.add_paragraph("原项目经历")
            paragraph.runs[0].bold = True
            document.save(source)

            blocks = extract_paragraph_blocks(source)
            target = next(block for block in blocks if block["text"] == "原项目经历")

            result = apply_paragraph_replacements(
                source_path=source,
                output_path=output,
                replacements=[
                    {
                        "block_id": target["block_id"],
                        "original_text": target["text"],
                        "new_text": "围绕内容运营目标完成项目方案设计、协作推进和结果复盘。",
                    }
                ],
            )

            self.assertEqual(result["applied"], 1)
            rewritten = Document(output)
            texts = [paragraph.text for paragraph in rewritten.paragraphs]
            self.assertIn(
                "围绕内容运营目标完成项目方案设计、协作推进和结果复盘。",
                texts,
            )

    def test_apply_paragraph_removals(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            source = directory / "resume.docx"
            output = directory / "tailored.docx"

            document = Document()
            document.add_paragraph("原项目经历")
            document.add_paragraph("无关校园经历")
            document.save(source)

            target = next(
                block
                for block in extract_paragraph_blocks(source)
                if block["text"] == "无关校园经历"
            )
            result = apply_paragraph_replacements(
                source_path=source,
                output_path=output,
                replacements=[],
                removals=[
                    {
                        "block_id": target["block_id"],
                        "original_text": target["text"],
                        "reason": "与目标岗位无关。",
                    }
                ],
            )

            self.assertEqual(result["removed"], 1)
            rewritten = Document(output)
            texts = [paragraph.text for paragraph in rewritten.paragraphs]
            self.assertNotIn("无关校园经历", texts)
            self.assertIn("原项目经历", texts)

    def test_tailored_resume_agent_returns_replacements(self) -> None:
        agent = TailoredResumeAgent(llm_client=FakeDeepSeekClient())
        result = agent.run(
            {
                "resume": {"skills": ["Python"]},
                "job_description": {"keywords": ["内容运营"]},
                "resume_blocks": [{"block_id": 1, "text": "原项目经历"}],
                "allow_skill_boosting": True,
            }
        )
        self.assertEqual(result["target_keywords"], ["内容运营", "项目管理"])
        self.assertEqual(len(result["replacements"]), 1)
        self.assertEqual(len(result["removals"]), 1)
        self.assertEqual(result["skill_boosts"][0]["skill"], "数据分析")

    def test_tailored_resume_agent_disables_skill_boosting(self) -> None:
        agent = TailoredResumeAgent(llm_client=FakeDeepSeekClient())
        result = agent.run(
            {
                "resume": {"skills": ["Python"]},
                "job_description": {"keywords": ["内容运营"]},
                "resume_blocks": [{"block_id": 1, "text": "原项目经历"}],
                "allow_skill_boosting": False,
            }
        )

        self.assertEqual(result["skill_boosts"], [])


if __name__ == "__main__":
    unittest.main()

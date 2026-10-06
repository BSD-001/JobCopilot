"""浏览器离线验收服务，仅使用虚构数据和临时历史目录。"""

from pathlib import Path
from tempfile import TemporaryDirectory
import sys

import uvicorn
from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from tests.web_fixtures import WebFakeClient
from webapp import create_app


if __name__ == "__main__":
    assets = PROJECT_ROOT / "output" / "playwright"
    assets.mkdir(parents=True, exist_ok=True)
    sample = Document()
    sample.add_paragraph("示例求职者")
    sample.add_paragraph("制作并交付3份项目说明。")
    sample.save(assets / "sample-resume.docx")
    with TemporaryDirectory(prefix="jobcopilot-browser-test-") as directory:
        application = create_app(Path(directory) / "history", WebFakeClient)
        uvicorn.run(application, host="127.0.0.1", port=8503, log_level="warning")

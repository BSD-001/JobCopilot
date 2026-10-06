"""验证四字能力与事实边界，不使用用户个人案例。"""

from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from docx import Document

from agents.resume_optimizer import ResumeOptimizerAgent
from agents.tailored_resume_agent import TailoredResumeAgent
from core.prompt_templates import AGENT_INSTRUCTIONS, OUTPUT_SCHEMAS
from core.resume_writing import check_experience_text, check_resume_text
from main import run_pipeline_result
from tests.web_fixtures import SAMPLE_JD, SAMPLE_RESUME, WebFakeClient


class ResumeWritingTests(TestCase):
    def test_pending_items_accept_strings_objects_and_empty_values(self):
        class PendingClient:
            def __init__(self, pending):
                self.pending = pending

            def generate_json(self, **kwargs):
                return {"pending_items": self.pending, "optimized_projects": [], "replacements": [], "skill_boosts": []}

        for agent_type in (ResumeOptimizerAgent, TailoredResumeAgent):
            for pending, expected in (
                (None, []),
                (" 补充项目反馈 ", ["补充项目反馈"]),
                ({"项目": "需求记录整理", "待补充": "使用反馈"}, ["项目：需求记录整理；待补充：使用反馈"]),
                ([None, "", "补充反馈", "补充反馈", {"项目": "需求记录整理", "待补充": ["结果口径", "人员范围"]}], ["补充反馈", '项目：需求记录整理；待补充：["结果口径", "人员范围"]']),
            ):
                with self.subTest(agent=agent_type.__name__, pending=pending):
                    result = agent_type(PendingClient(pending)).run({"resume_blocks": []})
                    self.assertEqual(result["pending_items"], expected)

    def test_object_pending_items_complete_docx_pipeline(self):
        class ObjectPendingClient(WebFakeClient):
            def generate_json(self, system_prompt, user_prompt, **kwargs):
                data = super().generate_json(system_prompt, user_prompt, **kwargs)
                if "pending_items" in data:
                    data["pending_items"] = [{"项目": "需求记录整理", "待补充": "使用反馈"}]
                return data

        with TemporaryDirectory() as directory:
            source = Path(directory) / "resume.docx"
            output = Path(directory) / "tailored.docx"
            document = Document()
            document.add_paragraph("示例求职者")
            document.add_paragraph("制作并交付3份项目说明。")
            document.save(source)
            stages = []
            result = run_pipeline_result(SAMPLE_RESUME, SAMPLE_JD, ObjectPendingClient(), resume_file_path=source, tailored_resume_output=output, on_progress=stages.append)
            self.assertEqual(result.analysis_data["pending_items"], ["项目：需求记录整理；待补充：使用反馈"])
            self.assertIn("待补充：使用反馈", result.report_text)
            self.assertTrue(output.exists())
            self.assertNotIn("使用反馈", "\n".join(paragraph.text for paragraph in Document(output).paragraphs))
            expected_stages = ["读取Word模板段落", "等待模型生成定制简历内容", "将定制内容回填到Word模板", "生成线上投递招呼语", "整理分析结果", "汇总完整分析报告"]
            self.assertEqual([stage for stage in stages if stage in expected_stages], expected_stages)

    def test_four_chinese_characters_and_adjacent_labels(self):
        self.assertEqual(check_experience_text("【文档交付】完成说明。\n【需求核对】记录反馈。"), [])
        self.assertTrue(check_experience_text("【文档】完成说明。"))
        self.assertTrue(check_experience_text("【文档交付】完成说明。\n【文档交付】记录反馈。"))

    def test_placeholder_and_new_numbers_are_rejected(self):
        self.assertTrue(check_resume_text("完成[待补充：数量]份文档。", "完成文档。"))
        self.assertTrue(check_resume_text("完成30份文档。", "完成3份文档。"))
        self.assertFalse(check_resume_text("完成3份文档。", "完成3份文档。"))

    def test_invalid_word_replacement_is_not_applied(self):
        class InvalidClient:
            def generate_json(self, **kwargs):
                return {"replacements": [{"block_id": 0, "new_text": "【文档交付】完成[待补充：数量]份报告。", "kind": "experience"}], "pending_items": [], "skill_boosts": []}
        result = TailoredResumeAgent(InvalidClient()).run({"resume_blocks": [{"text": "整理报告。"}]})
        self.assertEqual(result["replacements"], [])
        self.assertTrue(result["writing_warnings"])
        self.assertTrue(result["pending_items"])

    def test_common_rules_have_no_personal_template_constraints(self):
        for name in ("resume_optimizer", "tailored_resume_agent"):
            prompt = AGENT_INSTRUCTIONS[name]
            self.assertIn("四个汉字", prompt)
            self.assertIn("pending_items", prompt)
            self.assertNotIn("固定两页", prompt)
            self.assertNotIn("个人经历版", prompt)
            self.assertIn("pending_items", OUTPUT_SCHEMAS[name])

    def test_progress_and_structured_results_preserve_canonical_advice(self):
        stages = []
        result = run_pipeline_result(SAMPLE_RESUME, SAMPLE_JD, WebFakeClient(), on_progress=stages.append)
        self.assertIn("解析简历", stages)
        self.assertIn("生成面试问题", stages)
        self.assertIn("汇总完整分析报告", stages)
        self.assertIn("【文档交付】", result.report_text)
        self.assertIn("待补充信息", result.report_text)
        self.assertNotIn("模型扩写内容", result.report_text)
        self.assertEqual(result.analysis_data["match"]["score"], 73)

    def test_missing_report_section_still_preserves_advice(self):
        class MissingSectionClient(WebFakeClient):
            def generate_text(self, **kwargs):
                return "# 报告\n\n## 线上投递打招呼语\n\n需要替换的招呼语。"
        result = run_pipeline_result(SAMPLE_RESUME, SAMPLE_JD, MissingSectionClient())
        self.assertIn("【文档交付】", result.report_text)
        self.assertIn("待补充信息", result.report_text)
        self.assertNotIn("需要替换的招呼语", result.report_text)
        self.assertTrue(result.report_text.rstrip().endswith(result.analysis_data["hr_greeting"]))

    def test_disabled_or_unsafe_skill_boost_cannot_enter_word(self):
        class BoostClient:
            def generate_json(self, **kwargs):
                return {"replacements": [{"block_id": 0, "new_text": "熟练SQL", "kind": "other"}], "skill_boosts": [{"skill": "SQL", "suggested_wording": "熟练SQL"}]}
        for enabled in (True, False):
            result = TailoredResumeAgent(BoostClient()).run({"resume_blocks": [{"text": "文档整理"}], "allow_skill_boosting": enabled})
            self.assertEqual(result["skill_boosts"], [])
            self.assertEqual(result["replacements"], [])
            self.assertTrue(result["writing_warnings"])

    def test_skill_boost_retains_learning_wording_and_limit(self):
        class BoostClient:
            def generate_json(self, **kwargs):
                return {"replacements": [{"block_id": 0, "new_text": "文档整理；正在学习SQL", "kind": "other"}], "skill_boosts": [{"skill": skill, "suggested_wording": f"正在学习{skill}"} for skill in ("SQL", "Python", "Axure", "Figma")]}
        result = TailoredResumeAgent(BoostClient()).run({"resume_blocks": [{"text": "文档整理"}], "allow_skill_boosting": True})
        self.assertEqual(len(result["skill_boosts"]), 3)
        self.assertEqual(result["replacements"][0]["new_text"], "文档整理；正在学习SQL")

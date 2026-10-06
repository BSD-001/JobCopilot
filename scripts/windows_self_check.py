"""便携包离线自检：不调用真实模型，不访问用户历史。"""

import json
from pathlib import Path
import sys
from tempfile import TemporaryDirectory

from docx import Document
from pypdf import PdfWriter
import uvicorn

from core.document_reader import read_document
from core.history_store import load_history_record, save_history_record
from core.llm_client import DeepSeekClient
from core.resume_template import apply_paragraph_replacements, extract_paragraph_blocks
from core.runtime_paths import default_history_root
from webapp import PROJECT_ROOT, create_app, report_html


def run_self_check(result_path: Path) -> int:
    result = {"success": False, "frozen": bool(getattr(sys, "frozen", False)), "model_requests": 0}
    try:
        assert all((PROJECT_ROOT / "web" / name).is_file() for name in ("index.html", "styles.css", "app.js"))
        result["web_assets"] = True
        assert "<table>" in report_html("|列|\n|---|\n|值|")
        assert "<pre><code>" in report_html("```\n示例\n```")
        result["markdown_extensions"] = True
        with TemporaryDirectory(prefix="jobcopilot-package-check-") as directory:
            root = Path(directory)
            source, output = root / "sample.docx", root / "result.docx"
            document = Document()
            document.add_paragraph("示例求职者")
            document.add_paragraph("制作3份说明。")
            document.save(source)
            block = extract_paragraph_blocks(source)[1]
            apply_paragraph_replacements(source, output, [{"block_id": block["block_id"], "original_text": block["text"], "new_text": "【文档交付】制作3份说明。"}])
            assert "【文档交付】" in read_document(output)
            result["word_roundtrip"] = True
            pdf = PdfWriter()
            pdf.add_blank_page(width=100, height=100)
            pdf_path = root / "blank.pdf"
            with pdf_path.open("wb") as stream:
                pdf.write(stream)
            assert read_document(pdf_path) == ""
            result["pdf_reader"] = True
            record = save_history_record("示例岗位", "# 自检报告", output, root / "history", {"pending_items": ["示例反馈"]})
            assert load_history_record(record["record_id"], root / "history")["analysis_data"]["pending_items"] == ["示例反馈"]
            result["history_roundtrip"] = True
            application = create_app(root / "history")
            uvicorn.Config(application, loop="asyncio", http="h11", ws="none").load()
            application.state.executor.shutdown(wait=True)
            result["http_protocol"] = True
        client = DeepSeekClient(api_key="offline-package-check-not-a-real-key")
        client.client.close()
        result["llm_client_initialization"] = True
        if result["frozen"]:
            assert default_history_root() != PROJECT_ROOT / "output" / "history"
        result["history_outside_bundle"] = True
        result["success"] = True
    except Exception as error:
        result["exception_type"] = type(error).__name__
    result_path.parent.mkdir(parents=True, exist_ok=True)
    result_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))
    return 0 if result["success"] else 1

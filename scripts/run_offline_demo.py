"""不访问网络的完整流水线演示。"""

from __future__ import annotations

import sys
import re
from pathlib import Path
from typing import Any

from docx import Document


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.document_reader import read_document
from main import run_pipeline


DEMO_REPORT = """# JobCopilot 求职分析报告

> 演示说明：本报告由离线模拟客户端生成，用于展示完整流水线。它不是 DeepSeek 真实模型输出。

## 一、匹配度分析

综合评分：78 / 100

匹配点：

1. 简历中有活动策划、宣传执行和团队协作经历。
2. 简历体现了人工智能专业背景和人工智能工具使用经验。
3. 内容生产、项目推进和沟通协作与目标岗位方向一致。

缺失点：

1. 缺少明确的数据分析项目结果。
2. 缺少用户调研和产品需求文档相关经历。
3. 缺少可量化的内容运营结果。

## 二、简历优化建议

1. 社团嘉年华活动：突出“策划、跨部门协调、宣传物料产出和参与人数”，明确说明人工智能工具用于哪些具体环节。
2. 暑期社会实践：补充团队人数、服务地点、交付物数量和投稿渠道，避免只描述职责。
3. 人工智能工具实践项目：把使用的提示词方法、人工审校流程和最终作品整理成可展示成果。

## 三、面试预测题

1. 你为什么从人工智能专业转向产品运营方向？
2. 你如何使用人工智能工具完成一个真实项目？
3. 人工智能生成的内容出现错误时，你如何检查和修正？
4. 你如何判断一个内容是否适合目标用户？
5. 讲一个你协调多人完成项目的经历。
6. 你如何理解人工智能产品运营和传统运营的区别？
7. 你如何设计一次内容增长实验？
8. 项目时间不足时，你如何安排优先级？
9. 你如何衡量一次活动的效果？
10. 未来三十天你准备提升什么能力？

## 四、三十天学习计划

第一周：学习用户调研、岗位分析和需求整理，完成一份目标岗位分析。

第二周：学习内容运营、数据指标和复盘方法，完成一篇宣传内容拆解。

第三周：学习提示词设计、结构化输出和人工核验，完成一个可演示的人工智能工作流。

第四周：整理项目作品集，模拟面试并复盘，完成一版针对目标岗位的定制简历。

## 五、项目成果摘要

本次演示已经跑通以下流程：

1. 简历解析
2. 岗位描述解析
3. 匹配度分析
4. 简历优化
5. 面试问题预测
6. 三十天学习计划
7. 最终报告生成

## 线上投递打招呼语

您好，我关注到贵公司的人工智能产品运营岗位。我具备人工智能专业背景，也在真实项目中用 ChatGPT、Claude 和 Midjourney 等工具完成过文案、海报与内容策划。我很愿意从理解用户需求、推进内容和持续复盘做起，希望有机会进一步沟通。
"""


class OfflineDemoClient:
    """模拟 DeepSeek 客户端，专门用于离线演示。"""

    def generate_json(
        self,
        system_prompt: str,
        user_prompt: str,
        **kwargs: Any,
    ) -> dict[str, Any]:
        if "请解析简历文本" in user_prompt:
            return {
                "education": [
                    {
                        "school": "示例大学",
                        "major": "人工智能",
                        "degree": "本科",
                        "start_date": "2023.09",
                        "end_date": "2027.06",
                        "details": ["人工智能导论", "机器学习基础"],
                    }
                ],
                "projects": [
                    {
                        "name": "社团嘉年华",
                        "role": "总策划",
                        "start_date": "",
                        "end_date": "",
                        "description": "负责活动策划、宣传和现场统筹。",
                        "actions": ["使用人工智能工具生成文案和海报初稿"],
                        "results": ["完成活动宣传物料"],
                    }
                ],
                "skills": ["ChatGPT", "Midjourney", "Excel"],
                "experiences": [],
            }
        if "请解析岗位描述" in user_prompt:
            return {
                "job_title": "AI产品运营",
                "responsibilities": ["内容策划", "内容生产", "数据复盘"],
                "requirements": ["理解人工智能产品", "沟通协作", "项目推进"],
                "keywords": ["人工智能产品", "内容运营", "用户反馈"],
            }
        if "匹配度评分" in user_prompt:
            return {
                "score": 78,
                "matched_points": ["活动策划", "人工智能工具使用", "沟通协作"],
                "missing_points": ["数据分析", "用户调研", "内容运营数据"],
                "reason": "方向匹配，但缺少数据结果和产品需求文档经历。",
            }
        if "重写简历中的项目经历" in user_prompt:
            return {
                "optimized_projects": [
                    {"title": "活动策划", "content": "优化后的活动策划描述"},
                    {"title": "社会实践", "content": "优化后的社会实践描述"},
                    {"title": "人工智能工具实践", "content": "优化后的人工智能工具实践描述"},
                ]
            }
        if "高频面试问题" in user_prompt:
            return {
                "questions": [
                    {"question": f"演示面试题 {index}", "answer": "演示回答", "focus": "岗位匹配"}
                    for index in range(1, 11)
                ]
            }
        if "resume_blocks" in user_prompt:
            block_ids = [
                int(value)
                for value in re.findall(r'"block_id":\s*(\d+)', user_prompt)
            ][:2]
            return {
                "target_keywords": ["内容运营", "项目推进", "数据分析"],
                "summary": "针对岗位强化内容运营、项目推进和数据分析表达。",
                "replacements": [
                    {
                        "block_id": block_ids[1],
                        "original_text": "",
                        "new_text": "围绕内容运营目标完成项目方案设计、协作推进和结果复盘。",
                        "reason": "优先呈现岗位关注的关键词。",
                    }
                ],
                "removals": [
                    {
                        "block_id": block_ids[0],
                        "original_text": "无关校园经历",
                        "reason": "与目标岗位关联较弱。",
                    }
                ],
                "skill_boosts": [
                    {
                        "skill": "数据分析",
                        "suggested_wording": "正在学习数据分析基础",
                        "reason": "岗位高频要求。",
                    }
                ],
            }
        if "生成一段可直接发送给招聘平台 HR" in user_prompt:
            return {
                "greeting": "您好，我应聘AI产品运营。我有活动策划和AI内容生产经验，方便看下简历吗？"
            }
        return {
            "weeks": [
                {
                    "week": week,
                    "goal": f"第 {week} 周目标",
                    "tasks": ["完成学习任务", "整理成果"],
                    "deliverable": "学习成果",
                }
                for week in range(1, 5)
            ],
            "summary": "完成三十天学习计划。",
        }

    def generate_text(
        self,
        system_prompt: str,
        user_prompt: str,
        **kwargs: Any,
    ) -> str:
        return DEMO_REPORT


def main() -> int:
    resume_template_path = PROJECT_ROOT / "output" / "demo-resume-template.docx"
    resume_template_path.parent.mkdir(parents=True, exist_ok=True)
    template_document = Document()
    template_document.add_paragraph("无关校园经历")
    template_document.add_paragraph("原项目经历")
    template_document.save(resume_template_path)

    resume_text = read_document(resume_template_path)
    job_description_text = read_document(PROJECT_ROOT / "data" / "jd.txt")
    tailored_resume_path = PROJECT_ROOT / "output" / "demo-tailored-resume.docx"
    report_text = run_pipeline(
        resume_text=resume_text,
        job_description_text=job_description_text,
        llm_client=OfflineDemoClient(),
        resume_file_path=resume_template_path,
        tailored_resume_output=tailored_resume_path,
    )
    output_path = PROJECT_ROOT / "output" / "demo-report.md"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(report_text, encoding="utf-8")
    print(f"演示报告已生成：{output_path}")
    print(f"岗位定制简历已生成：{tailored_resume_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
